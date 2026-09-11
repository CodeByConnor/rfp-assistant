"""Export a reviewed run.

The xlsx export writes answers back into the buyer's own workbook instead of
producing a new one, because the filled-in workbook is the deliverable. The
workbook is loaded normally -- never `data_only`, which would flatten every
formula to a value -- and only vendor-owned cells are written, so the buyer's
scoring formulas, dropdowns, and formatting survive the round trip.

Three rules, each tested:

- Only approved rows are written. Unreviewed model output never reaches the
  customer's file; pending rows are left blank.
- A citation to an internal-only document is never written, whatever the row
  says.
- A run made under the internal role cannot be exported without an explicit
  override, because its answers may quote material the customer must not see.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .classify import Classification
from .knowledge import PUBLIC
from .models import Requirement
from .report import render_markdown
from .review import APPROVED, Run
from .workbook import vendor_columns


class ExportError(ValueError):
    pass


@dataclass
class ExportResult:
    written: int
    pending: int
    citations_withheld: int


def _public_citation(row) -> str | None:
    return row.citation if row.citation and row.citation_access == PUBLIC else None


def export_xlsx(
    run: Run,
    source: str | Path,
    out_path: str | Path,
    *,
    allow_internal: bool = False,
) -> ExportResult:
    if run.role != PUBLIC and not allow_internal:
        raise ExportError(
            "this run was made with internal-only documents visible and may contain "
            "material that must not reach the customer; it cannot be exported"
        )

    approved = {r.rid: r for r in run.rows if r.status == APPROVED}
    pending = len(run.rows) - len(approved)
    withheld = sum(1 for r in approved.values() if r.citation and _public_citation(r) is None)

    source = Path(source)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        written = _round_trip(source, Path(out_path), approved)
    else:
        written = _fresh_workbook(run, Path(out_path), approved)

    return ExportResult(written=written, pending=pending, citations_withheld=withheld)


def _round_trip(source: Path, out_path: Path, approved: dict) -> int:
    wb = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    written = 0
    try:
        for ws in wb.worksheets:
            found = vendor_columns(ws)
            if not found:
                continue
            header_row, cols = found
            for row in range(header_row + 1, ws.max_row + 1):
                rid = ws.cell(row=row, column=cols["id"]).value
                rid = str(rid).strip() if rid is not None else ""
                answer = approved.get(rid)
                if answer is None:
                    continue
                if "compliance" in cols and answer.compliance_level:
                    ws.cell(row=row, column=cols["compliance"]).value = answer.compliance_level
                if "response" in cols:
                    ws.cell(row=row, column=cols["response"]).value = answer.answer
                if "reference" in cols:
                    ws.cell(row=row, column=cols["reference"]).value = _public_citation(answer)
                written += 1
        wb.save(out_path)
    finally:
        wb.close()
    return written


def _fresh_workbook(run: Run, out_path: Path, approved: dict) -> int:
    """For pdf and markdown sources there is no buyer workbook to fill, so the
    export mirrors the RFP as a response table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Response"
    headers = ["Req ID", "Requirement", "Compliance Level", "Response", "Source"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(name="Arial", bold=True)
    for col, width in zip("ABCDE", (10, 70, 18, 70, 28)):
        ws.column_dimensions[col].width = width

    written = 0
    for i, row in enumerate(run.rows, start=2):
        ws.cell(row=i, column=1, value=row.rid)
        ws.cell(row=i, column=2, value=row.text)
        answer = approved.get(row.rid)
        if answer is not None:
            ws.cell(row=i, column=3, value=answer.compliance_level)
            ws.cell(row=i, column=4, value=answer.answer)
            ws.cell(row=i, column=5, value=_public_citation(answer))
            written += 1
        for col in range(1, 6):
            cell = ws.cell(row=i, column=col)
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return written


def export_markdown(run: Run) -> str:
    requirements = {
        r.rid: Requirement(
            rid=r.rid,
            text=r.text,
            section=r.section,
            section_title=r.section_title,
            priority=r.priority,
            weight=r.weight,
            source_format=run.source_format,
            locator=r.locator,
        )
        for r in run.rows
    }
    results = [
        Classification(
            rid=r.rid,
            verdict=r.verdict,
            answer=r.answer,
            citation=r.citation,
            supporting_quote=r.supporting_quote,
            overridden_from=r.overridden_from,
            override_reason=r.override_reason,
        )
        for r in run.rows
    ]
    counts = run.counts()
    header = (
        f"_Source: {run.source_name} · mode: {run.mode} · "
        f"{counts['approved']} of {counts['total']} answers approved._\n\n"
    )
    body = render_markdown(requirements, results, title=f"Gap Report — {run.source_name}")
    first_break = body.index("\n\n") + 2
    return body[:first_break] + header + body[first_break:]
