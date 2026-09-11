"""Workbook column discovery, shared by run creation and export.

The parser only needs the id and requirement columns. Writing answers back
needs the vendor-owned columns too, and needs to know which requirement tabs
carry no compliance column at all -- narrative tabs like Vendor Profile, where
"Standard" or "Not Supported" is meaningless.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .parsers.xlsx_parser import _find_header

COMPLIANCE_HEADERS = ("compliance level", "compliance")
RESPONSE_HEADERS = ("vendor comments", "vendor response", "response", "comments")
REFERENCE_HEADERS = ("supporting doc reference", "supporting documentation", "reference document")


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def vendor_columns(ws) -> tuple[int, dict[str, int]] | None:
    """Header row and column indexes for id plus any vendor-owned columns."""
    found = _find_header(ws)
    if not found:
        return None
    header_row, cols = found
    labels = {
        _norm(ws.cell(row=header_row, column=col).value): col
        for col in range(1, (ws.max_column or 0) + 1)
    }
    labels.pop("", None)
    out = {"id": cols["id"]}
    for key, names in (
        ("compliance", COMPLIANCE_HEADERS),
        ("response", RESPONSE_HEADERS),
        ("reference", REFERENCE_HEADERS),
    ):
        col = next((labels[name] for name in names if name in labels), None)
        if col is not None:
            out[key] = col
    return header_row, out


def narrative_tabs(path: str | Path) -> set[str]:
    """Requirement tabs that have no compliance column."""
    wb = load_workbook(path)
    try:
        tabs = set()
        for ws in wb.worksheets:
            found = vendor_columns(ws)
            if found and "compliance" not in found[1]:
                tabs.add(ws.title)
        return tabs
    finally:
        wb.close()
