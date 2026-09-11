"""Review workflow, export, and web app tests. All offline."""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "fixtures"))

from rfp_assistant.classify import NEEDS_INPUT, classify_all  # noqa: E402
from rfp_assistant.export import ExportError, export_markdown, export_xlsx  # noqa: E402
from rfp_assistant.knowledge import INTERNAL, PUBLIC, load_knowledge_base  # noqa: E402
from rfp_assistant.llm import ReplayClient, StubClient  # noqa: E402
from rfp_assistant.parsers import parse  # noqa: E402
from rfp_assistant.retrieval import BM25Retriever  # noqa: E402
from rfp_assistant.review import (  # noqa: E402
    APPROVED,
    PENDING,
    NotFound,
    ReviewError,
    create_run,
    list_runs,
    load_run,
    save_run,
    source_path,
    update_row,
)
from rfp_assistant.web import app as web  # noqa: E402

KB = ROOT / "docs" / "knowledge-base"
XLSX = ROOT / "fixtures" / "rfp-alderwood-retail.xlsx"
MD = ROOT / "fixtures" / "rfp-alderwood-retail.md"
GOLD = ROOT / "fixtures" / "gold-answers.json"


@pytest.fixture(scope="module")
def retriever():
    return BM25Retriever(load_knowledge_base(KB))


def make_run(store, retriever, source=XLSX, *, role=PUBLIC, client=None):
    requirements = parse(source)
    results, _ = classify_all(
        requirements, retriever, client or ReplayClient(GOLD), role=role
    )
    return create_run(
        store, source, source.name, requirements, results, role=role, mode="replay"
    )


@pytest.fixture
def run(tmp_path, retriever):
    return make_run(tmp_path, retriever)


def first(run, test):
    return next(r for r in run.rows if test(r))


# ---------------------------------------------------------------- review rules

def test_unambiguous_verdicts_prefill_a_compliance_level(run):
    row = first(run, lambda r: r.verdict == "Yes" and not r.narrative)
    assert row.compliance_level == "Standard"
    row = first(run, lambda r: r.verdict == "No" and not r.narrative)
    assert row.compliance_level == "Not Supported"


def test_partial_is_not_prefilled_because_it_is_a_human_call(run):
    """Partial could be Configuration, Customization, or Third-Party, and the
    buyer scores those differently."""
    row = first(run, lambda r: r.verdict == "Partial" and not r.narrative)
    assert row.compliance_level is None
    with pytest.raises(ReviewError, match="compliance level"):
        update_row(run, row.rid, status=APPROVED)
    update_row(run, row.rid, compliance_level="Configuration", status=APPROVED)
    assert row.status == APPROVED


def test_needs_input_cannot_be_approved_with_placeholder_text(run):
    row = first(run, lambda r: r.verdict == NEEDS_INPUT and not r.narrative)
    with pytest.raises(ReviewError, match="written answer"):
        update_row(run, row.rid, compliance_level="Not Supported", status=APPROVED)
    update_row(
        run, row.rid, answer="Answered by the account team.",
        compliance_level="Not Supported", status=APPROVED,
    )
    assert row.status == APPROVED and row.edited


def test_editing_an_approved_answer_returns_it_to_pending(run):
    row = first(run, lambda r: r.verdict == "Yes" and not r.narrative)
    update_row(run, row.rid, status=APPROVED)
    update_row(run, row.rid, answer=row.answer + " Revised.")
    assert row.status == PENDING


def test_narrative_rows_take_no_compliance_level(run):
    row = first(run, lambda r: r.narrative)
    assert row.compliance_level is None
    with pytest.raises(ReviewError, match="narrative"):
        update_row(run, row.rid, compliance_level="Standard")


def test_narrative_rows_approve_without_a_compliance_level(run):
    row = first(run, lambda r: r.narrative and r.verdict != NEEDS_INPUT)
    update_row(run, row.rid, status=APPROVED)
    assert row.status == APPROVED and row.compliance_level is None


def test_narrative_rows_come_from_the_tab_without_a_compliance_column(run):
    narrative = {r.locator.split("!")[0] for r in run.rows if r.narrative}
    assert narrative == {"Vendor Profile"}


@pytest.mark.parametrize(
    "kwargs, message",
    [
        ({"compliance_level": "Mostly"}, "unknown compliance level"),
        ({"status": "done"}, "unknown status"),
        ({"answer": "   "}, "must not be empty"),
        ({"answer": "x" * 4001}, "limited"),
    ],
)
def test_invalid_updates_are_rejected(run, kwargs, message):
    row = first(run, lambda r: not r.narrative)
    with pytest.raises(ReviewError, match=message):
        update_row(run, row.rid, **kwargs)


def test_unknown_requirement_is_not_found(run):
    with pytest.raises(NotFound):
        update_row(run, "99.99.99", status=APPROVED)


# ------------------------------------------------------------------ storage

def test_run_round_trips_through_disk(tmp_path, run):
    row = first(run, lambda r: r.verdict == "Yes" and not r.narrative)
    update_row(run, row.rid, status=APPROVED)
    save_run(tmp_path, run)
    loaded = load_run(tmp_path, run.id)
    assert loaded.row(row.rid).status == APPROVED
    assert len(loaded.rows) == 248
    assert [r.id for r in list_runs(tmp_path)] == [run.id]


@pytest.mark.parametrize("bad", ["../etc", "..", "", "ABCDEF0123456789", "0123456789abcdef/x"])
def test_run_ids_cannot_reach_the_filesystem(tmp_path, bad):
    with pytest.raises(NotFound):
        load_run(tmp_path, bad)


def test_the_uploaded_source_is_kept_with_the_run(tmp_path, run):
    assert source_path(tmp_path, run).read_bytes() == XLSX.read_bytes()


# ------------------------------------------------------------------- export

@pytest.fixture
def exported(tmp_path, run):
    yes = first(run, lambda r: r.verdict == "Yes" and not r.narrative)
    partial = first(run, lambda r: r.verdict == "Partial" and not r.narrative)
    update_row(run, yes.rid, status=APPROVED)
    update_row(run, partial.rid, compliance_level="Customization", status=APPROVED)
    out = tmp_path / "out.xlsx"
    result = export_xlsx(run, source_path(tmp_path, run), out)
    return result, load_workbook(out), yes, partial


def _row_values(wb, rid):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            if row[0].value == rid:
                return ws.title, [cell.value for cell in row]
    raise AssertionError(f"{rid} not found")


def test_only_approved_rows_are_written(exported, run):
    result, wb, yes, partial = exported
    assert result.written == 2
    assert result.pending == len(run.rows) - 2

    _, values = _row_values(wb, yes.rid)
    assert values[4] == "Standard" and values[5] == yes.answer

    _, values = _row_values(wb, partial.rid)
    assert values[4] == "Customization"

    untouched = first(run, lambda r: r.status == PENDING and not r.narrative)
    _, values = _row_values(wb, untouched.rid)
    assert values[4] is None and values[5] is None


def test_export_preserves_the_buyers_formulas_and_dropdowns(exported):
    """The workbook is the deliverable. Flattening the buyer's scoring model
    to static values would break the file they send to their evaluators."""
    _, wb, _, _ = exported
    assert str(wb["Scoring Summary"]["B16"].value).startswith("=SUM(")
    assert str(wb["Security"]["I11"].value).startswith("=IFERROR(")
    assert wb["Functional"].data_validations.dataValidation


def test_export_does_not_touch_the_example_row(exported):
    _, wb, _, _ = exported
    example = next(r for r in wb["Functional"].iter_rows(values_only=True) if r[0] == "EXAMPLE")
    assert example[4] == "Configuration"  # the buyer's own illustrative value


def test_internal_citations_are_never_written(tmp_path, run):
    row = first(run, lambda r: r.verdict == "Yes" and not r.narrative)
    row.citation, row.citation_access = "pricing-packaging.md", INTERNAL
    update_row(run, row.rid, status=APPROVED)
    out = tmp_path / "out.xlsx"
    result = export_xlsx(run, source_path(tmp_path, run), out)
    assert result.citations_withheld == 1
    _, values = _row_values(load_workbook(out), row.rid)
    assert "pricing-packaging.md" not in [v for v in values if isinstance(v, str)]


def test_unknown_citation_provenance_fails_closed(run):
    """A citation whose chunk was not among the evidence is treated as internal."""
    for row in run.rows:
        if row.citation:
            assert row.citation_access in {PUBLIC, INTERNAL}
            if row.override_reason == "citation not in evidence":
                assert row.citation_access == INTERNAL


def test_internal_role_runs_cannot_be_exported(tmp_path, retriever):
    run = make_run(tmp_path, retriever, role=INTERNAL, client=StubClient())
    with pytest.raises(ExportError):
        export_xlsx(run, source_path(tmp_path, run), tmp_path / "out.xlsx")


def test_pdf_and_markdown_sources_export_a_fresh_response_sheet(tmp_path, retriever):
    run = make_run(tmp_path, retriever, source=MD)
    row = first(run, lambda r: r.verdict == "Yes")
    update_row(run, row.rid, status=APPROVED)
    out = tmp_path / "out.xlsx"
    result = export_xlsx(run, source_path(tmp_path, run), out)
    ws = load_workbook(out)["Response"]
    assert result.written == 1
    assert ws.max_row == 249  # header + every requirement, filled or not


def test_markdown_report_includes_review_progress(run):
    md = export_markdown(run)
    assert md.startswith("# Gap Report")
    assert f"0 of {len(run.rows)} answers approved" in md


# ------------------------------------------------------------ replay client

def test_replay_client_returns_labelled_verdicts():
    client = ReplayClient(GOLD)
    response = client.complete(
        "system", "Requirement 5.1.4 [Must]\ntext\n\nEvidence:\n[evidence 1] security-whitepaper.md | x\n"
        "Meridian is not HIPAA compliant and does not execute Business Associate Agreements.\n",
    )
    assert response.data["verdict"] == "No"


def test_replay_needs_input_carries_no_citation():
    client = ReplayClient(GOLD)
    response = client.complete("s", "Requirement 8.3.1\ntext\n\nEvidence:\n[evidence 1] a.md | h\nsome long line of evidence text here for the stub\n")
    assert response.data["verdict"] == NEEDS_INPUT
    assert response.data["citation"] == ""


# --------------------------------------------------------------------- app

@pytest.fixture
def client(tmp_path):
    app = web.create_app(
        store_dir=tmp_path / "runs",
        kb_dir=KB,
        client_factory=lambda: ReplayClient(GOLD),
        mode="replay",
        sample_path=XLSX,
    )
    return TestClient(app)


def test_config_reports_offline_mode(client):
    config = client.get("/api/config").json()
    assert config["live"] is False and config["mode"] == "replay"
    assert "Standard" in config["compliance_levels"]


def test_sample_run_then_review_then_export(client):
    run = client.post("/api/sample").json()
    assert run["counts"]["total"] == 248
    row = next(r for r in run["rows"] if r["verdict"] == "Yes" and not r["narrative"])

    res = client.patch(f"/api/runs/{run['id']}/rows/{row['rid']}", json={"status": "approved"})
    assert res.status_code == 200 and res.json()["counts"]["approved"] == 1

    res = client.get(f"/api/runs/{run['id']}/export.xlsx")
    assert res.status_code == 200
    assert res.headers["x-export-written"] == "1"
    assert res.headers["content-disposition"].endswith('rfp-alderwood-retail-response.xlsx"')


def test_rejected_approval_surfaces_the_reason(client):
    run = client.post("/api/sample").json()
    row = next(r for r in run["rows"] if r["verdict"] == "Partial" and not r["narrative"])
    res = client.patch(f"/api/runs/{run['id']}/rows/{row['rid']}", json={"status": "approved"})
    assert res.status_code == 400
    assert "compliance level" in res.json()["detail"]


def test_upload_accepts_each_supported_format(client):
    for source in (XLSX, MD, ROOT / "fixtures" / "rfp-alderwood-retail.pdf"):
        res = client.post("/api/runs", files={"file": (source.name, source.read_bytes())})
        assert res.status_code == 201, source.name
        assert res.json()["counts"]["total"] == 248


def test_upload_filename_is_display_only(client, tmp_path):
    res = client.post(
        "/api/runs", files={"file": ("../../../evil.md", MD.read_bytes())}
    )
    assert res.status_code == 201
    assert res.json()["source_name"] == "evil.md"
    assert not (tmp_path.parent / "evil.md").exists()


@pytest.mark.parametrize(
    "name, data, status",
    [
        ("malware.exe", b"MZ", 415),
        ("empty.md", b"", 400),
        ("nothing.md", b"# No numbered requirements here\n", 422),
        ("broken.xlsx", b"not really a workbook", 422),
    ],
)
def test_bad_uploads_are_rejected(client, name, data, status):
    assert client.post("/api/runs", files={"file": (name, data)}).status_code == status


def test_oversized_upload_is_rejected(client, monkeypatch):
    monkeypatch.setattr(web, "MAX_UPLOAD_BYTES", 100)
    res = client.post("/api/runs", files={"file": ("big.md", b"x" * 500)})
    assert res.status_code == 413


@pytest.mark.parametrize("run_id", ["not-a-run", "0123456789abcdef"])
def test_unknown_runs_are_404(client, run_id):
    assert client.get(f"/api/runs/{run_id}").status_code == 404


def test_the_browser_cannot_switch_to_live_mode(client):
    """There is no endpoint that changes the client. Live mode is decided when
    the server starts, after a terminal confirmation."""
    routes = {getattr(r, "path", "") for r in client.app.routes}
    assert not any("live" in path or "config" in path and path != "/api/config" for path in routes)
    assert client.post("/api/config", json={"live": True}).status_code == 405


def test_limit_caps_requirements_per_upload(tmp_path):
    app = web.create_app(
        store_dir=tmp_path, kb_dir=KB, client_factory=StubClient,
        mode="stub", limit=10, sample_path=XLSX,
    )
    run = TestClient(app).post("/api/sample").json()
    assert run["counts"]["total"] == 10


def test_page_is_served(client):
    res = client.get("/")
    assert res.status_code == 200
    assert "RFP Assistant" in res.text
